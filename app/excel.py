"""Planilla Excel para descargar desde el panel.

- "Lista de seguridad": una fila por persona confirmada, en orden alfabetico y lista para imprimir.
- "Invitaciones": una fila por grupo, con todos los datos (uso interno).

Todo texto se guarda como texto: un nombre como "=1+1" no se ejecuta como formula y un telefono
"+54 9 ..." no se convierte en numero.
"""
import io
import unicodedata
from datetime import datetime
from typing import Iterable, List, Sequence, Tuple

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties
from openpyxl.worksheet.worksheet import Worksheet

from app import config
from app.models import Invitacion, Tipo

_linea = Side(style="thin", color="BFBFBF")
BORDE = Border(left=_linea, right=_linea, top=_linea, bottom=_linea)
FUENTE_ENCABEZADO = Font(bold=True, color="FFFFFF")
FONDO_ENCABEZADO = PatternFill("solid", fgColor="5E6D58")
ARRIBA = Alignment(vertical="top", wrap_text=True)


def _clave_orden(texto: str) -> str:
    sin_tildes = unicodedata.normalize("NFKD", texto or "").encode("ascii", "ignore").decode()
    return sin_tildes.casefold()


def _texto_encabezado(texto: str) -> str:
    return texto.replace("&", "&&")  # en encabezados de pagina "&" es un codigo de Excel


def _fila(hoja: Worksheet, valores: Sequence, con_borde: bool = True) -> None:
    hoja.append(list(valores))
    for celda in hoja[hoja.max_row]:
        if isinstance(celda.value, str):
            celda.data_type = "s"  # openpyxl toma como formula todo texto que empieza con "="
        if con_borde:
            celda.border = BORDE
            celda.alignment = ARRIBA


def _preparar(hoja: Worksheet, titulo: str, columnas: List[Tuple[str, int]], horizontal: bool) -> None:
    """Encabezado con estilo, fila fija, titulos repetidos al imprimir y ajuste al ancho de la hoja."""
    hoja.title = titulo
    _fila(hoja, [nombre for nombre, _ in columnas])
    for i, (celda, (_, ancho)) in enumerate(zip(hoja[1], columnas), start=1):
        celda.font = FUENTE_ENCABEZADO
        celda.fill = FONDO_ENCABEZADO
        hoja.column_dimensions[get_column_letter(i)].width = ancho
    hoja.freeze_panes = "A2"
    hoja.print_title_rows = "1:1"
    hoja.page_setup.paperSize = hoja.PAPERSIZE_A4
    hoja.page_setup.orientation = "landscape" if horizontal else "portrait"
    hoja.page_setup.fitToWidth = 1
    hoja.page_setup.fitToHeight = 0
    hoja.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    hoja.oddHeader.center.text = _texto_encabezado(f"{config.NOVIA} & {config.NOVIO} · {titulo}")
    hoja.oddHeader.right.text = "Generado " + datetime.now(config.TZ).strftime("%d/%m/%Y %H:%M")
    hoja.oddFooter.center.text = "Página &P de &N"


def _lista_seguridad(hoja: Worksheet, invitaciones: Iterable[Invitacion]) -> None:
    personas = sorted(
        (
            (g.nombre or "(sin nombre)", inv.nombre_grupo, g.tipo, g.restriccion or "")
            for inv in invitaciones
            for g in inv.invitados
            if g.asiste is True
        ),
        key=lambda p: (_clave_orden(p[0]), _clave_orden(p[1])),
    )
    _preparar(hoja, "Lista de seguridad", [
        ("N°", 6), ("Nombre y apellido", 34), ("Grupo / familia", 28), ("Tipo", 9),
        ("Restricción alimentaria", 30), ("Llegó", 8),
    ], horizontal=False)
    for numero, (nombre, grupo, tipo, restriccion) in enumerate(personas, start=1):
        _fila(hoja, [numero, nombre, grupo, "Niño/a" if tipo == Tipo.nino else "Adulto", restriccion, ""])
    if personas:
        hoja.auto_filter.ref = f"A1:F{hoja.max_row}"
    ninos = sum(1 for p in personas if p[2] == Tipo.nino)
    hoja.append([])
    _fila(hoja, ["", f"Total confirmados: {len(personas)} ({len(personas) - ninos} adultos, {ninos} niños)"],
          con_borde=False)
    hoja.cell(row=hoja.max_row, column=2).font = Font(bold=True)


def _invitaciones(hoja: Worksheet, invitaciones: Iterable[Invitacion]) -> None:
    _preparar(hoja, "Invitaciones", [
        ("Código", 9), ("Grupo / familia", 28), ("Formato", 9), ("Adultos", 8), ("Niños", 7),
        ("Estado", 11), ("Confirmados", 12), ("Ingresados", 11), ("Teléfono", 17), ("Email", 26),
        ("Invitados", 34), ("Restricciones", 30), ("Mensaje", 40), ("Link", 38),
    ], horizontal=True)
    for inv in invitaciones:
        invitados = []
        for g in inv.invitados:
            marca = "asiste" if g.asiste else ("no asiste" if g.asiste is False else "sin respuesta")
            invitados.append(f"{g.nombre or '(sin nombre)'} — {marca}")
        restricciones = [f"{g.nombre or '(sin nombre)'}: {g.restriccion}" for g in inv.invitados if g.restriccion]
        _fila(hoja, [
            inv.codigo, inv.nombre_grupo, "Física" if inv.tarjeta_fisica else "Virtual",
            inv.cupo_adultos, inv.cupo_ninos, inv.estado.value.capitalize(),
            inv.confirmados, inv.ingresados, inv.telefono or "", inv.email or "",
            "\n".join(invitados), "\n".join(restricciones), inv.mensaje or "",
            f"{config.BASE_URL}/i/{inv.codigo}",
        ])
    if hoja.max_row > 1:
        hoja.auto_filter.ref = f"A1:N{hoja.max_row}"


def generar(invitaciones: Sequence[Invitacion]) -> bytes:
    libro = Workbook()
    _lista_seguridad(libro.active, invitaciones)
    _invitaciones(libro.create_sheet(), invitaciones)
    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()
